"""Converts the report story (same flowables as the PDF) into Lab7_Report.docx with python-docx."""
import os, sys, re, html
os.environ["LAB7_DOCX"] = "1"
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE); sys.path.insert(0, os.path.dirname(HERE))
import build_report as br
from reportlab.platypus import Paragraph, Table, Image, KeepTogether, Spacer, PageBreak, NextPageTemplate
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TOK = re.compile(r"(<[^>]+>)")


def add_runs(par, markup, size=None, bold_all=False):
    markup = markup.replace("<br/>", "\x00").replace("<br />", "\x00")
    st = dict(b=bold_all, i=False, sub=False, sup=False, mono=False)
    for tok in TOK.split(markup):
        if not tok:
            continue
        if tok.startswith("<"):
            t = tok.strip("<>/ ").split()[0].lower() if tok.strip("<>/ ") else ""
            on = not tok.startswith("</")
            if t in ("b", "strong"): st["b"] = on or bold_all
            elif t in ("i", "em"): st["i"] = on
            elif t in ("sub", "sup"): st[t] = on
            elif t == "font": st["mono"] = on and "Courier" in tok
            continue
        for k, piece in enumerate(tok.split("\x00")):
            if k:
                par.add_run().add_break()
            txt = html.unescape(piece)
            if not txt:
                continue
            r = par.add_run(txt)
            r.bold = st["b"] or None; r.italic = st["i"] or None
            r.font.subscript = st["sub"] or None; r.font.superscript = st["sup"] or None
            if st["mono"]: r.font.name = "Courier New"
            if size: r.font.size = Pt(size)


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)


class Conv:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]; s.page_width, s.page_height = Mm(210), Mm(297)
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, m, Mm(14))
        self.d.styles["Normal"].font.name = "Calibri"; self.d.styles["Normal"].font.size = Pt(9.5)
        self.width_mm = 182
        self.landscape = False

    def section(self, land):
        s = self.d.add_section(WD_SECTION.NEW_PAGE)
        if land:
            s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Mm(297), Mm(210); self.width_mm = 270
        else:
            s.orientation = WD_ORIENT.PORTRAIT; s.page_width, s.page_height = Mm(210), Mm(297); self.width_mm = 182
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, m, Mm(10 if land else 14))
        self.landscape = land

    def para(self, f, container=None, size=None):
        name = f.style.name
        d = container or self.d
        if name in ("h1", "h2", "h3") and container is None:
            p = d.add_heading(level={"h1": 1, "h2": 2, "h3": 3}[name]); add_runs(p, f.text); return p
        if name == "title" and container is None:
            p = d.add_heading(level=0); add_runs(p, f.text); return p
        p = d.add_paragraph()
        sz = size or {"small": 8, "tiny": 7, "cap": 8, "cell": 8, "cellb": 8, "note": 9}.get(name)
        add_runs(p, f.text, sz, bold_all=(name == "cellb"))
        if name == "cap":
            for r in p.runs: r.italic = True
        if container is not None:
            p.paragraph_format.space_after = Pt(1)
        return p

    def image(self, im, container_par=None):
        path = im.filename
        w = min(im.drawWidth / 2.8346, self.width_mm - 4)          # points -> mm
        if container_par is None:
            container_par = self.d.add_paragraph()
        container_par.add_run().add_picture(path, width=Mm(w))

    def table(self, t):
        rows = t._cellvalues; nrows = len(rows); ncols = max(len(r) for r in rows)
        tb = self.d.add_table(rows=nrows, cols=ncols); tb.style = "Table Grid"; tb.autofit = False
        widths = getattr(t, "_colWidths", None)
        tot = sum(w for w in widths if w) if widths and all(widths) else None
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                cell = tb.cell(i, j)
                if widths and tot:
                    cell.width = Mm(widths[j] / 2.8346 * self.width_mm / (tot / 2.8346))
                cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
                items = c if isinstance(c, list) else [c]
                got = False
                for it in items:
                    if isinstance(it, Paragraph):
                        self.para(it, container=cell, size=7 if self.landscape else 8); got = True
                    elif isinstance(it, Image):
                        p = cell.add_paragraph(); self.image(it, p); got = True
                    elif isinstance(it, Table):
                        # nested figure rows
                        for rr in it._cellvalues:
                            for cc in rr:
                                for x in (cc if isinstance(cc, list) else [cc]):
                                    if isinstance(x, Image):
                                        p = cell.add_paragraph(); self.image(x, p); got = True
                if not got:
                    cell.add_paragraph("")
                if i == 0:
                    shade(cell, "DFE7F3")
        # spans (family header rows)
        for cmd in getattr(t, "_bkgrndcmds", []):
            pass
        self.d.add_paragraph()

    def flow(self, f):
        if isinstance(f, Paragraph): self.para(f)
        elif isinstance(f, Table):
            rows = f._cellvalues
            if len(rows) == 1 and all(isinstance(x, Image) for x in rows[0]):
                p = self.d.add_paragraph()
                for x in rows[0]:
                    p.add_run().add_picture(x.filename, width=Mm(min(x.drawWidth / 2.8346, self.width_mm / len(rows[0]) - 2)))
            else:
                self.table(f)
        elif isinstance(f, Image): self.image(f)
        elif isinstance(f, KeepTogether):
            for x in f._content: self.flow(x)
        elif isinstance(f, PageBreak):
            self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif isinstance(f, NextPageTemplate):
            self.pending = f.action[1]
        elif isinstance(f, Spacer):
            pass


def main():
    story = br.main()
    c = Conv(); c.pending = None
    for f in story:
        if isinstance(f, NextPageTemplate):
            c.pending = f.action[1]; continue
        if isinstance(f, PageBreak):
            if c.pending:
                c.section(c.pending == "landscape"); c.pending = None
            else:
                c.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            continue
        c.flow(f)
    out = os.path.join(os.path.dirname(HERE), "Lab7_Report.docx")
    c.d.save(out); print("wrote", out)


if __name__ == "__main__":
    main()
